"""Stage 6: convert the feature selection report to Word (.docx).

What it does
    Builds a .docx file from the same Markdown file the PDF is built from, so
    the text and every number are identical in the three formats. Nothing is
    recomputed here. The Markdown is rendered to HTML (same step as the PDF)
    and each HTML element is written as its Word equivalent: headings,
    paragraphs, bullet lists, bordered tables and embedded figures. Text
    marked as code in the Markdown (feature names, configuration names, file
    paths) is written in a monospaced font.

Input
    report/feature_selection_report.md (or the _es.md version) and the
    figures it references in report/figures/

Output
    report/feature_selection_report.docx      English (default)
    report/feature_selection_report_es.docx   Spanish, with --lang es

How to run (after 06b_report.py with the same language)
    .venv/bin/python src/06d_report_docx.py
    .venv/bin/python src/06d_report_docx.py --lang es

Needs the packages `markdown` and `python-docx` in the venv. The fonts used
(Arial, Courier New) exist in Word; LibreOffice replaces them with
metric-compatible fonts, so the layout is the same in both.
"""
import datetime
import re
import sys
import xml.etree.ElementTree as ET

import markdown
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

from config import (REPORT_DATE, REPORT_DIR, REPORT_DOCX_FILE,
                    REPORT_DOCX_FILE_ES, REPORT_FILE, REPORT_FILE_ES)

BODY_FONT = "Arial"
CODE_FONT = "Courier New"
INK = RGBColor(0x1F, 0x23, 0x28)
# Printable width of an A4 page with the margins set in new_document().
CONTENT_WIDTH_MM = 178
SINGLE_FIGURE_WIDTH_MM = 165
# Two figures share a page when each is at most this tall.
PAIRED_FIGURE_HEIGHT_MM = 104


def set_font(run, name, size=None):
    """Set the font of a run for Latin and for other scripts.

    Word keeps a separate font setting for non-Latin text; setting both
    avoids a fallback font on accented characters.
    """
    run.font.name = name
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    for attribute in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attribute), name)
    if size:
        run.font.size = Pt(size)


def new_document():
    """Create an A4 document with the report's base styles."""
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.left_margin = section.right_margin = Mm(16)
    section.top_margin = section.bottom_margin = Mm(18)

    sizes = {"Normal": 10, "Heading 1": 17, "Heading 2": 13, "Heading 3": 11,
             "List Bullet": 10}
    for name, size in sizes.items():
        style = doc.styles[name]
        style.font.name = BODY_FONT
        style.font.size = Pt(size)
        style.font.color.rgb = INK
        fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
        for attribute in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            fonts.set(qn(attribute), BODY_FONT)
        # The built-in heading styles point to a theme font, which would
        # override the font name set above.
        for attribute in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme",
                          "w:eastAsiaTheme"):
            fonts.attrib.pop(qn(attribute), None)
        if name.startswith("Heading"):
            style.font.bold = True
    doc.styles["Normal"].paragraph_format.space_after = Pt(6)
    add_page_number(section)
    return doc


def add_page_number(section):
    """Put a centered page number in the footer."""
    paragraph = section.footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    set_font(run, BODY_FONT, 8)
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            element = OxmlElement("w:fldChar")
            element.set(qn("w:fldCharType"), kind)
        else:
            element = OxmlElement("w:instrText")
            element.set(qn("xml:space"), "preserve")
            element.text = text
        run._element.append(element)


def add_inline(paragraph, node, size=None, bold=False, code=False):
    """Write the text of ``node`` and its inline children into ``paragraph``.

    Handles bold (<strong>), italics (<em>) and code (<code>); code is the
    part that must come out monospaced.
    """
    def write(text, bold, code, italic=False):
        if not text:
            return
        # Markdown source lines are wrapped (and indented inside lists); a
        # line break inside a paragraph is just one space.
        run = paragraph.add_run(re.sub(r"[ \t]*\n[ \t]*", " ", text))
        run.bold = bold or None
        run.italic = italic or None
        if code:
            set_font(run, CODE_FONT, (size or 10) - 0.5)
        elif size:
            set_font(run, BODY_FONT, size)

    write(node.text, bold, code)
    for child in node:
        if child.tag == "strong":
            add_inline(paragraph, child, size, True, code)
        elif child.tag == "code":
            add_inline(paragraph, child, size, bold, True)
        elif child.tag == "em":
            write("".join(child.itertext()), bold, code, italic=True)
        else:
            add_inline(paragraph, child, size, bold, code)
        write(child.tail, bold, code)


def add_figures(doc, node):
    """Embed the image(s) of a paragraph.

    One image takes almost the full page width. Two images in the same
    Markdown paragraph are meant to share a page: both are limited in height
    and the first one is kept on the same page as the second.
    """
    images = node.findall("img")
    paired = len(images) > 1
    for position, image in enumerate(images):
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(4)
        paragraph.paragraph_format.keep_with_next = position < len(images) - 1
        path = str(REPORT_DIR / image.get("src"))
        run = paragraph.add_run()
        if paired:
            picture = run.add_picture(path, height=Mm(PAIRED_FIGURE_HEIGHT_MM))
        else:
            picture = run.add_picture(path, width=Mm(SINGLE_FIGURE_WIDTH_MM))
        # Alternative text, for screen readers and for checking the file.
        picture._inline.docPr.set("descr", image.get("alt", ""))


def shade(cell, color):
    """Fill a table cell with a background color."""
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), color)
    cell._element.get_or_add_tcPr().append(shading)


def column_weight(cells):
    """Relative width a table column needs.

    A column must be at least as wide as its longest word, because a word
    (for example a configuration name) cannot be broken without becoming hard
    to read. Beyond that, columns with long sentences get more room, up to a
    limit, so a text column does not squeeze the numeric ones.
    """
    longest_word, longest_text = 0, 0
    for cell in cells:
        text = "".join(cell.itertext())
        # Monospaced characters are wider than the proportional font.
        factor = 1.3 if cell.find("code") is not None else 1.0
        words = text.split()
        if words:
            longest_word = max(longest_word,
                               max(len(w) for w in words) * factor)
        longest_text = max(longest_text, len(text))
    return max(longest_word + 2, min(longest_text, 60) * 0.6)


def add_table(doc, node):
    """Write an HTML table as a bordered Word table that fits the page.

    Column widths follow the longest text of each column, scaled so the table
    is exactly as wide as the printable area: Word then wraps long cells
    instead of pushing the table past the margin.
    """
    rows = [[cell for cell in row] for row in node.iter("tr")]
    n_cols = max(len(row) for row in rows)
    # The widest tables (nine columns) need a smaller font to fit.
    size = 7 if n_cols >= 7 else 8

    table = doc.add_table(rows=len(rows), cols=n_cols)
    table.style = "Table Grid"  # the built-in style with all borders
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    weights = [column_weight([row[c] for row in rows if c < len(row)])
               for c in range(n_cols)]
    widths = [CONTENT_WIDTH_MM * w / sum(weights) for w in weights]
    # The column grid is what LibreOffice reads; Word reads the cell widths
    # set below. Both are written so the two programs agree.
    for c, width in enumerate(widths):
        table.columns[c].width = Mm(width)

    for r, row in enumerate(rows):
        header = row[0].tag == "th"
        for c, cell_node in enumerate(row):
            cell = table.cell(r, c)
            cell.width = Mm(widths[c])
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            add_inline(paragraph, cell_node, size=size, bold=header)
            if header:
                shade(cell, "F3F4F6")
        properties = table.rows[r]._tr.get_or_add_trPr()
        # A row is never split between two pages.
        properties.append(OxmlElement("w:cantSplit"))
        if header:
            # Repeat the header row when the table continues on a new page.
            properties.append(OxmlElement("w:tblHeader"))
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def build(doc, root):
    """Write every block element of the rendered report into ``doc``."""
    for node in root:
        if node.tag in ("h1", "h2", "h3"):
            heading = doc.add_heading(level=int(node.tag[1]))
            add_inline(heading, node)
            text = "".join(node.itertext())
            # The appendix starts on a new page, as in the PDF.
            if text.startswith(("Appendix", "Anexo")):
                heading.paragraph_format.page_break_before = True
        elif node.tag == "p" and node.find("img") is not None:
            add_figures(doc, node)
        elif node.tag == "p":
            add_inline(doc.add_paragraph(), node)
        elif node.tag == "ul":
            for item in node.findall("li"):
                add_inline(doc.add_paragraph(style="List Bullet"), item)
        elif node.tag == "table":
            add_table(doc, node)
        else:
            raise ValueError(f"unsupported element in the report: {node.tag}")


def main():
    """Render the Markdown report and save it as .docx."""
    lang = "en"
    if "--lang" in sys.argv:
        lang = sys.argv[sys.argv.index("--lang") + 1]
    source_file, target_file = {
        "en": (REPORT_FILE, REPORT_DOCX_FILE),
        "es": (REPORT_FILE_ES, REPORT_DOCX_FILE_ES)}[lang]

    source = (REPORT_DIR / source_file).read_text(encoding="utf-8")
    body = markdown.markdown(source, extensions=["tables", "sane_lists"])
    # The HTML produced by the markdown package is well-formed, so it can be
    # read as XML once it has a single root element.
    root = ET.fromstring(f"<root>{body}</root>")

    doc = new_document()
    build(doc, root)
    # Fixed metadata so the file does not change between runs.
    stamp = datetime.datetime.fromisoformat(REPORT_DATE)
    doc.core_properties.created = doc.core_properties.modified = stamp
    doc.core_properties.author = "Laura Camila Herrera Murillo"
    doc.core_properties.title = "".join(root.find("h1").itertext())
    target = REPORT_DIR / target_file
    doc.save(target)
    print(f"Written: {target} ({target.stat().st_size / 1024:.0f} KB)")


if __name__ == "__main__":
    main()
